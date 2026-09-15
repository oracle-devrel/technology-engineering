"""Create in-memory JSON, CSV, and Excel audit packs.

Author: Ali Ottoman
"""

from __future__ import annotations

import io
import json
import re
from datetime import datetime, timezone

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ExtractionRun

INK = "151715"


def _spreadsheet_safe(value):
    """Prevent user or model text from becoming an Excel/CSV formula."""

    if not isinstance(value, str):
        return value
    cleaned = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", value)[:32_767]
    visible = cleaned.lstrip(" \t\r\n")
    return f"'{cleaned}" if visible.startswith(("=", "+", "-", "@")) else cleaned


def runs_to_json(runs: list[ExtractionRun]) -> bytes:
    payload = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "generated_by": "Mizan Invoice Intelligence · OCI Generative AI",
        "documents": [run.model_dump(mode="json") for run in runs],
    }
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")


def runs_to_csv(runs: list[ExtractionRun]) -> bytes:
    rows: list[dict] = []
    for run in runs:
        invoice = run.extraction.invoice
        for item in run.extraction.line_items:
            rows.append(
                {
                    "file": run.filename,
                    "document_type": run.extraction.document_type,
                    "invoice_number": invoice.invoice_number,
                    "invoice_date": invoice.invoice_date,
                    "supplier": run.extraction.supplier.name,
                    "customer": run.extraction.customer.name,
                    "currency": invoice.currency,
                    "validation_status": (
                        "blocked"
                        if run.validation.errors
                        else "review"
                        if run.validation.reviews
                        else "passed"
                    ),
                    "approval_status": run.approval_status,
                    "approved_by": run.approved_by,
                    "approved_at": run.approved_at,
                    **item.model_dump(mode="json"),
                }
            )
    buffer = io.StringIO()
    safe_rows = [
        {key: _spreadsheet_safe(value) for key, value in row.items()} for row in rows
    ]
    pd.DataFrame(safe_rows).to_csv(buffer, index=False)
    return buffer.getvalue().encode("utf-8-sig")


def runs_to_xlsx(runs: list[ExtractionRun]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Invoices"
    _append_rows(summary, _summary_rows(runs))

    lines = workbook.create_sheet("Line Items")
    _append_rows(lines, _line_rows(runs))

    taxes = workbook.create_sheet("Tax")
    _append_rows(taxes, _tax_rows(runs))

    audit = workbook.create_sheet("Validation")
    _append_rows(audit, _audit_rows(runs))

    evidence = workbook.create_sheet("Evidence")
    _append_rows(evidence, _evidence_rows(runs))

    custom = workbook.create_sheet("Custom Fields")
    _append_rows(custom, _custom_rows(runs))

    changes = workbook.create_sheet("Review Changes")
    _append_rows(changes, _change_rows(runs))

    passes = workbook.create_sheet("Pass History")
    _append_rows(passes, _pass_rows(runs))

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _summary_rows(runs: list[ExtractionRun]) -> list[list]:
    rows = [
        [
            "File",
            "Document type",
            "Invoice #",
            "Supplier",
            "Customer",
            "Invoice date",
            "Due date",
            "PO #",
            "Currency",
            "Subtotal",
            "Tax",
            "Total",
            "Balance due",
            "Review score",
            "Required for approval",
            "Status",
            "Approved by",
            "Approved at",
        ]
    ]
    for run in runs:
        data = run.extraction
        rows.append(
            [
                run.filename,
                data.document_type,
                data.invoice.invoice_number,
                data.supplier.name,
                data.customer.name,
                data.invoice.invoice_date,
                data.invoice.due_date,
                data.invoice.purchase_order,
                data.invoice.currency,
                data.totals.subtotal,
                data.totals.total_tax,
                data.totals.total_amount,
                data.totals.balance_due,
                run.validation.score,
                ", ".join(run.required_field_paths),
                (
                    "Approved"
                    if run.approval_status == "approved"
                    else "Blocked"
                    if run.validation.errors
                    else "Pending review"
                ),
                run.approved_by,
                run.approved_at,
            ]
        )
    return rows


def _line_rows(runs: list[ExtractionRun]) -> list[list]:
    rows = [
        [
            "File",
            "Invoice #",
            "Line",
            "Description (original)",
            "Description (English)",
            "Product code",
            "Quantity",
            "Unit",
            "Unit price",
            "Discount",
            "Tax rate",
            "Amount",
            "Page",
            "Confidence",
        ]
    ]
    for run in runs:
        for item in run.extraction.line_items:
            rows.append(
                [
                    run.filename,
                    run.extraction.invoice.invoice_number,
                    item.line_number,
                    item.description_original,
                    item.description_english,
                    item.product_code,
                    item.quantity,
                    item.unit,
                    item.unit_price,
                    item.discount,
                    item.tax_rate,
                    item.amount,
                    item.source_page,
                    item.confidence,
                ]
            )
    return rows


def _tax_rows(runs: list[ExtractionRun]) -> list[list]:
    rows = [["File", "Invoice #", "Tax", "Rate", "Taxable amount", "Tax amount"]]
    for run in runs:
        for tax in run.extraction.tax_lines:
            rows.append(
                [
                    run.filename,
                    run.extraction.invoice.invoice_number,
                    tax.name,
                    tax.rate,
                    tax.taxable_amount,
                    tax.tax_amount,
                ]
            )
    return rows


def _audit_rows(runs: list[ExtractionRun]) -> list[list]:
    rows = [["File", "Check", "Status", "Message", "Fields"]]
    for run in runs:
        for check in run.validation.checks:
            rows.append(
                [
                    run.filename,
                    check.label,
                    check.severity.upper(),
                    check.message,
                    ", ".join(check.field_paths),
                ]
            )
    return rows


def _evidence_rows(runs: list[ExtractionRun]) -> list[list]:
    rows = [["File", "Field", "Page", "Source quote", "Confidence"]]
    for run in runs:
        for item in run.extraction.evidence:
            rows.append(
                [
                    run.filename,
                    item.field_path,
                    item.source_page,
                    item.quote,
                    item.confidence,
                ]
            )
    return rows


def _custom_rows(runs: list[ExtractionRun]) -> list[list]:
    rows = [["File", "Invoice #", "Field", "Original", "English", "Page", "Confidence"]]
    for run in runs:
        for field in run.extraction.custom_fields:
            rows.append(
                [
                    run.filename,
                    run.extraction.invoice.invoice_number,
                    field.name,
                    field.value_original,
                    field.value_english,
                    field.source_page,
                    field.confidence,
                ]
            )
    return rows


def _change_rows(runs: list[ExtractionRun]) -> list[list]:
    rows = [["File", "Field", "Before", "After", "Changed at", "Evidence"]]
    for run in runs:
        for change in run.review_changes:
            rows.append(
                [
                    run.filename,
                    change.field_path,
                    change.before,
                    change.after,
                    change.changed_at,
                    change.evidence_disposition,
                ]
            )
    return rows


def _pass_rows(runs: list[ExtractionRun]) -> list[list]:
    rows = [
        [
            "File",
            "Pass",
            "Type",
            "Score",
            "Issues",
            "Changed fields",
            "Note",
            "Actor",
            "Recorded at",
        ]
    ]
    for run in runs:
        for item in run.passes:
            rows.append(
                [
                    run.filename,
                    item.pass_number,
                    item.kind,
                    item.score,
                    item.issues,
                    ", ".join(item.changes),
                    item.note,
                    item.actor,
                    item.recorded_at,
                ]
            )
    return rows


def _append_rows(sheet, rows: list[list]) -> None:
    for row in rows:
        sheet.append([_spreadsheet_safe(value) for value in row])
    if not rows:
        return
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, start=1):
        values = [len(str(cell.value or "")) for cell in column[:120]]
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(values + [10]) + 2, 48
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
